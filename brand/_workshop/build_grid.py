import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

NAVY="1F3352"; RED="A5322C"; YEL="FFF6D6"; LGREY="F1EFEA"; INK="1A1A1A"
navy_fill=PatternFill("solid",fgColor=NAVY)
red_fill=PatternFill("solid",fgColor=RED)
yel_fill=PatternFill("solid",fgColor=YEL)
grey_fill=PatternFill("solid",fgColor=LGREY)
whiteB=Font(name="Calibri",bold=True,color="FFFFFF",size=11)
navyB=Font(name="Calibri",bold=True,color=NAVY,size=11)
redB=Font(name="Calibri",bold=True,color="FFFFFF",size=12)
base=Font(name="Calibri",color=INK,size=11)
ital=Font(name="Calibri",italic=True,color="666666",size=10)
title=Font(name="Georgia",bold=True,color=NAVY,size=16)
sub=Font(name="Georgia",italic=True,color="666666",size=11)
thin=Side(style="thin",color="D9D5CE")
border=Border(left=thin,right=thin,top=thin,bottom=thin)
center=Alignment(horizontal="center",vertical="center",wrap_text=True)
left=Alignment(horizontal="left",vertical="center",wrap_text=True)

wb=openpyxl.Workbook()

# ============ SHEET 1: Règles & Poids ============
s=wb.active; s.title="Règles & Poids"
s["A1"]="K E B R A N E — Grille de décision  ·  Logo"; s["A1"].font=title
s["A2"]="Instrument verrouillé avant notation. On fixe les poids d'abord ; on score à froid ensuite."; s["A2"].font=sub
s["A4"]="Date de verrouillage des poids :"; s["B4"]="2026-07-29"
s["A5"]="Responsable de la décision :"; s["B5"]="Profi"
s["A6"]="Principe directeur :"; s["B6"]="Le sens naît du retrait. La barrière retirée est le dessin."
for r in (4,5,6):
    s[f"A{r}"].font=navyB; s[f"B{r}"].font=base; s[f"A{r}"].alignment=left; s[f"B{r}"].alignment=left

def band(cell,text):
    s[cell]=text; s[cell].font=redB; s[cell].fill=red_fill; s[cell].alignment=left

band("A8","POIDS DES CRITÈRES — verrouillés AVANT toute notation")
s.merge_cells("A8:C8")
hdr=["Critère","Poids (%)","Ce que ça mesure"]
for i,h in enumerate(hdr):
    c=s.cell(row=9,column=1+i,value=h); c.font=whiteB; c.fill=navy_fill; c.alignment=center; c.border=border
crit=[
 ("Intégrité soustractive",25,"Le sens naît-il vraiment d'un retrait, pas d'un ajout ?"),
 ("Distinction",20,"Ne ressemble à rien de connu ; pas de confusion avec une icône générique."),
 ("Mémorabilité",20,"Reconstructible de mémoire par un inconnu."),
 ("Extensibilité",20,"Tient en verrou de sous-marque et en icône d'app carrée."),
 ("Intemporalité",15,"Ne sera pas daté dans cinq ans."),
]
r=10
for name,w,desc in crit:
    s.cell(row=r,column=1,value=name).font=base
    wc=s.cell(row=r,column=2,value=w); wc.font=Font(name="Calibri",bold=True,color="0000FF"); wc.fill=yel_fill; wc.alignment=center
    s.cell(row=r,column=3,value=desc).font=base
    for col in (1,2,3):
        s.cell(row=r,column=col).border=border; s.cell(row=r,column=col).alignment=left
    s.cell(row=r,column=2).alignment=center
    r+=1
s.cell(row=r,column=1,value="TOTAL").font=navyB
tc=s.cell(row=r,column=2,value="=SUM(B10:B14)"); tc.font=navyB; tc.alignment=center; tc.border=border
s.cell(row=r,column=3,value="Doit être égal à 100.").font=ital
s.cell(row=r,column=1).border=border

band("A17","SEUILS DE VALIDATION (absolus, non relatifs)")
s.merge_cells("A17:C17")
s["A18"]="Score pondéré minimum (/100)"; s["B18"]=70
s["A19"]="Plancher par critère (/5)"; s["B19"]=3
s["A20"]="Filtres éliminatoires"; s["B20"]="Tous = Oui"
for rr in (18,19,20):
    s[f"A{rr}"].font=base; s[f"A{rr}"].alignment=left
    s[f"B{rr}"].alignment=center; s[f"B{rr}"].border=border
    if rr in (18,19):
        s[f"B{rr}"].font=Font(name="Calibri",bold=True,color="0000FF"); s[f"B{rr}"].fill=yel_fill
    else:
        s[f"B{rr}"].font=base

band("A22","ÉCHELLE DE NOTATION (0–5, à appliquer à froid)")
s.merge_cells("A22:C22")
ech=[("0","Absent / le critère n'est pas rempli."),
     ("1–2","Faible : présent mais fragile."),
     ("3","Seuil acceptable (plancher)."),
     ("4","Solide."),
     ("5","Exemplaire.")]
r=23
for lvl,desc in ech:
    s.cell(row=r,column=1,value=lvl).font=navyB; s.cell(row=r,column=1).alignment=center
    s.cell(row=r,column=2,value=desc).font=base; s.merge_cells(f"B{r}:C{r}"); s.cell(row=r,column=2).alignment=left
    r+=1

band("A29","RÈGLE D'ARRÊT")
s.merge_cells("A29:C29")
s["A30"]=("Aucune direction n'est déclarée gagnante avant que toutes aient passé les filtres éliminatoires. "
          "Une direction est VALIDÉE seulement si : elle passe tous les filtres, atteint le score minimum, "
          "et n'a aucun critère sous le plancher. À égalité de score entre directions validées, la plus simple l'emporte "
          "(arbitrage manuel, assumé par écrit). L'attachement se note, mais ne vote pas.")
s.merge_cells("A30:C31"); s["A30"].font=base; s["A30"].alignment=left
s.column_dimensions["A"].width=30; s.column_dimensions["B"].width=16; s.column_dimensions["C"].width=52

# ============ SHEET 2: Filtres éliminatoires ============
f=wb.create_sheet("Filtres éliminatoires")
f["A1"]="FILTRES ÉLIMINATOIRES — tout ou rien"; f["A1"].font=title
f["A2"]="Réponds Oui / Non. Une seule case « Non » écarte la direction, quel que soit son charme."; f["A2"].font=sub
cols=["Direction","Lisible à 16 px\nen 1 couleur","Reproductible\nnavy + réservé","Indépendant\ndu O","Neutralité\nculturelle","VERDICT"]
for i,h in enumerate(cols):
    c=f.cell(row=4,column=1+i,value=h); c.font=whiteB; c.fill=navy_fill; c.alignment=center; c.border=border
dirs=["A — Le A ouvert","B — La brèche","C — Le seuil","D — Les deux E"]
for j,d in enumerate(dirs):
    row=5+j
    f.cell(row=row,column=1,value=d).font=navyB; f.cell(row=row,column=1).alignment=left; f.cell(row=row,column=1).border=border
    for col in range(2,6):
        cc=f.cell(row=row,column=col); cc.fill=yel_fill; cc.alignment=center; cc.border=border; cc.font=base
    vc=f.cell(row=row,column=6,value=f'=IF(COUNTA(B{row}:E{row})<4,"À compléter",IF(COUNTIF(B{row}:E{row},"Oui")=4,"ADMISE","ÉCARTÉE"))')
    vc.font=navyB; vc.alignment=center; vc.border=border
# example row
er=9
f.cell(row=er,column=1,value="Exemple (à effacer)").font=ital
for col,val in zip(range(2,6),["Oui","Oui","Oui","Oui"]):
    cc=f.cell(row=er,column=col,value=val); cc.font=ital; cc.alignment=center; cc.fill=grey_fill; cc.border=border
f.cell(row=er,column=6,value=f'=IF(COUNTA(B{er}:E{er})<4,"À compléter",IF(COUNTIF(B{er}:E{er},"Oui")=4,"ADMISE","ÉCARTÉE"))').font=ital
f.cell(row=er,column=6).alignment=center
f["A11"]="Légende : cases jaunes = à remplir (Oui/Non). Le verdict se calcule tout seul."; f["A11"].font=ital
f.column_dimensions["A"].width=20
for col in "BCDE": f.column_dimensions[col].width=15
f.column_dimensions["F"].width=15

# ============ SHEET 3: Notation pondérée ============
n=wb.create_sheet("Notation pondérée")
n["A1"]="NOTATION PONDÉRÉE — à froid, après les filtres"; n["A1"].font=title
n["A2"]="Note chaque critère de 0 à 5 (cases jaunes). Le score, les contrôles et le verdict sont automatiques."; n["A2"].font=sub
heads=["Direction","Intégrité\nsoustractive","Distinction","Mémorabilité","Extensibilité","Intemporalité",
       "Score pondéré\n/100","Plancher\n(aucun < 3)","Éliminatoires\nOK ?","VALIDÉE ?","Attachement\n(non décisionnel)"]
for i,h in enumerate(heads):
    c=n.cell(row=4,column=1+i,value=h); c.font=whiteB; c.fill=navy_fill; c.alignment=center; c.border=border
# weights row (references sheet 1)
n.cell(row=5,column=1,value="Poids (%) →").font=navyB; n.cell(row=5,column=1).alignment=Alignment(horizontal="right")
wref=["B10","B11","B12","B13","B14"]
for i,ref in enumerate(wref):
    c=n.cell(row=5,column=2+i,value=f"='Règles & Poids'!{ref}"); c.font=Font(name="Calibri",bold=True,color=NAVY); c.alignment=center; c.border=border; c.fill=grey_fill
for col in range(7,12):
    n.cell(row=5,column=col).fill=grey_fill; n.cell(row=5,column=col).border=border
elim_rows={0:5,1:6,2:7,3:8}  # direction index -> row in Filtres sheet
dnames=["A — Le A ouvert","B — La brèche","C — Le seuil","D — Les deux E"]
for j,d in enumerate(dnames):
    row=6+j
    n.cell(row=row,column=1,value=d).font=navyB; n.cell(row=row,column=1).alignment=left; n.cell(row=row,column=1).border=border
    for col in range(2,7):
        cc=n.cell(row=row,column=col); cc.fill=yel_fill; cc.alignment=center; cc.border=border; cc.font=base
    fr=elim_rows[j]
    n.cell(row=row,column=7,value=f"=SUMPRODUCT(B{row}:F{row},$B$5:$F$5)/5").font=navyB
    n.cell(row=row,column=8,value=f'=IF(COUNTA(B{row}:F{row})<5,"—",IF(MIN(B{row}:F{row})>=\'Règles & Poids\'!$B$19,"Oui","Non"))').font=base
    n.cell(row=row,column=9,value=f"=IF('Filtres éliminatoires'!F{fr}=\"ADMISE\",\"Oui\",\"Non\")").font=base
    n.cell(row=row,column=10,value=f'=IF(AND(G{row}>=\'Règles & Poids\'!$B$18,H{row}="Oui",I{row}="Oui"),"VALIDÉE","non")').font=navyB
    ac=n.cell(row=row,column=11); ac.fill=yel_fill; ac.font=ital; ac.alignment=center; ac.border=border
    for col in (7,8,9,10):
        n.cell(row=row,column=col).alignment=center; n.cell(row=row,column=col).border=border
n.cell(row=11,column=1,value="Rappel : la colonne « Attachement » rend visible ce qu'on aime — elle n'entre dans aucun calcul.").font=ital
n.merge_cells("A11:K11")
n.column_dimensions["A"].width=20
for col in "BCDEF": n.column_dimensions[col].width=12
for col in ["G","H","I","J"]: n.column_dimensions[col].width=13
n.column_dimensions["K"].width=18

# ============ SHEET 4: Verdict ============
v=wb.create_sheet("Verdict & Registre")
v["A1"]="VERDICT & REGISTRE"; v["A1"].font=title
v["A2"]="Synthèse automatique. À égalité entre directions validées : la plus simple l'emporte (à trancher et à écrire ici)."; v["A2"].font=sub
vh=["Direction","Score pondéré /100","VALIDÉE ?"]
for i,h in enumerate(vh):
    c=v.cell(row=4,column=1+i,value=h); c.font=whiteB; c.fill=navy_fill; c.alignment=center; c.border=border
for j,d in enumerate(dnames):
    row=5+j; nr=6+j
    v.cell(row=row,column=1,value=d).font=navyB; v.cell(row=row,column=1).border=border; v.cell(row=row,column=1).alignment=left
    v.cell(row=row,column=2,value=f"='Notation pondérée'!G{nr}").font=base; v.cell(row=row,column=2).alignment=center; v.cell(row=row,column=2).border=border
    v.cell(row=row,column=3,value=f"='Notation pondérée'!J{nr}").font=navyB; v.cell(row=row,column=3).alignment=center; v.cell(row=row,column=3).border=border
v["A10"]="Décision finale :"; v["A10"].font=navyB
v["B10"]=""; v["B10"].fill=yel_fill; v["B10"].border=border
v["A11"]="Justification (par écrit) :"; v["A11"].font=navyB
v.merge_cells("B11:C12"); v["B11"].fill=yel_fill; v["B11"].border=border; v["B11"].alignment=left
v.column_dimensions["A"].width=22; v.column_dimensions["B"].width=20; v.column_dimensions["C"].width=22

wb.save("Grille-de-decision-Logo-Kebrane.xlsx")
print("saved")
